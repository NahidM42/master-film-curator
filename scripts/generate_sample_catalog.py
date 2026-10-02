"""Generate the public-safe 60-record synthetic Rev07-compatible workbook."""

from argparse import ArgumentParser
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAIN_HEADERS = [
    "Rank", "Master ID", "Title", "Year", "Director / Creator", "Genre",
    "Rev06 Channel", "Rev07 Channel", "Rev07 Taste Fit", "Rev07 Tier",
    "Rev07 Viewing Priority", "Score Delta", "Analytical Note",
    "Rev07 Calibration Reason",
]
FLAG_HEADERS = [
    "Master ID", "Title", "Year", "Director", "Genre", "Rev07 Channel",
    "Rev07 Taste", "Tier", "Viewing Priority", "Why / Note",
]
TOP_HEADERS = [
    "Rank", "Master ID", "Title", "Year", "Director", "Rev07 Channel",
    "Rev07 Taste", "Tier", "Key Theme", "Central Question",
    "Possible Episode Angle", "Analytical Note",
]
AUDIT_HEADERS = [
    "Master ID", "Title", "Year", "Director", "Rev06 Channel", "Rev07 Channel",
    "Delta", "Rev06 Taste", "Rev07 Taste", "Rev06 Tier", "Rev07 Tier",
    "Rev06 Priority", "Rev07 Priority", "Rev06 Reason",
    "Rev07 Calibration Reason",
]

TITLES = [
    ("Lanterns at Noon", 1966, "Mira Voss", "Drama"),
    ("Glass Orchard", 2011, "Leila Arden", "Drama, Mystery"),
    ("Northbound Silence", 2003, "Tomas Vale", "Drama"),
    ("Paper Constellation", 2012, "Inez Rowan", "Drama"),
    ("Silent Harbor", 1968, "Arman Kestrel", "Drama"),
    ("Silent Harbor", 2011, "Omar Keene", "Drama, Thriller"),
    ("The Witness Room", 1973, "Pavel Mora", "Drama"),
    ("Meridian House", 1973, "Mira Voss", "Drama, Miniseries"),
    ("Amber District", 1981, "Sofia Deren", "Crime, Drama"),
    ("The Geometry of Rain", 1987, "Jonas Wren", "Drama"),
    ("Second Horizon", 1990, "Nadia Rook", "Drama, Adventure"),
    ("Blue Static", 1994, "Elias North", "Drama, Sci-Fi"),
    ("The Empty Balcony", 1998, "Rina Sol", "Drama"),
    ("A Map of Ash", 2001, "Caleb Mire", "Drama, Mystery"),
    ("Winter Archive", 2004, "Tara Quill", "Documentary"),
    ("The Quiet Engine", 2006, "Marek Lune", "Drama"),
    ("Borrowed Weather", 2008, "Nora Flint", "Drama, Romance"),
    ("Night Index", 2010, "Soren Pike", "Thriller"),
    ("Mosaic Street", 2013, "Anya Crest", "Drama"),
    ("The Long Tuesday", 2014, "Felix Raye", "Comedy, Drama"),
    ("Copper Sky", 2015, "Mina Thorne", "Drama"),
    ("Static Garden", 2016, "Oren Bell", "Drama, Sci-Fi"),
    ("A Door Left Open", 2017, "Selene March", "Drama"),
    ("Blackwater Letters", 2018, "Dorian Pike", "Mystery"),
    ("The Fourth Window", 2019, "Juno Vale", "Drama"),
    ("Common Ground", 2020, "Lena Orr", "Drama"),
    ("Echoes of Tuesday", 2021, "Ilya Stone", "Drama"),
    ("The Small Republic", 2022, "Mara Fen", "Drama"),
    ("After the Brass Bell", 2023, "Theo Nair", "Drama"),
    ("Folding Light", 2024, "Ada Linn", "Drama"),
    ("Signal Orchard", 1976, "Rafi Dune", "Drama"),
    ("Red Thread Station", 1983, "Nika Vale", "Drama"),
    ("The Borrowed Shore", 1989, "Yara Moss", "Drama"),
    ("Noon Without Shadows", 1993, "Milo Crane", "Drama"),
    ("North Glass", 1997, "Elio Fern", "Drama"),
    ("Cartography of Dust", 2000, "Sara Lark", "Documentary"),
    ("The Last Semaphore", 2002, "Kian Rook", "Drama"),
    ("Orchard of Signals", 2005, "Vera Noll", "Drama"),
    ("A Minor Eclipse", 2007, "Aris Dawn", "Drama"),
    ("The Listening Field", 2009, "Maya Crest", "Drama"),
    ("Counterfactual City", 2012, "Noel Venn", "Drama, Sci-Fi"),
    ("The Slow Argument", 2014, "Lior Beck", "Drama"),
    ("Faint Republic", 2016, "Tessa Moor", "Drama"),
    ("The Memory Bridge", 2018, "Caro Flint", "Drama"),
    ("Rooms of Weather", 2020, "Dara Quill", "Drama, TV Series"),
    ("The Parallel Guest", 2021, "Sami North", "Drama"),
    ("Thirty Quiet Steps", 2022, "Lina Rook", "Drama"),
    ("The Orchard Trial", 2023, "Mira Sol", "Drama"),
    ("Stone Atlas", 2024, "Tomas Crest", "Drama"),
    ("Border of Morning", 2025, "Nadia Vale", "Drama"),
    ("Glass Orchard", 2001, "Jori Wren", "Drama"),
    ("The Witness Room", 2022, "Selene Mora", "Drama, Thriller"),
    ("Northbound Silence", 1992, "Daria Lune", "Drama"),
    ("Northbound Silence", 2024, "Pavel Orr", "Drama"),
    ("A Map of Ash", 2016, "Marek Fen", "Drama"),
    ("A Map of Ash", 2024, "Tara Pike", "Drama"),
    ("Open Circuit", 2017, "Lena Voss", "Documentary, TV Series"),
    ("Rooms of Weather: Field Notes", 2021, "Dara Quill", "Documentary, Reality-TV"),
    ("The Ninth Landing", 2019, "Inez Bell", "Drama"),
    ("Under a Paper Moon", 2025, "Arman Nair", "Drama"),
]

NOT_FOR_CHANNEL = {1010, 1019, 1028, 1037, 1046, 1053, 1057, 1060}
NOT_FOR_TASTE = {1012, 1021, 1028, 1034, 1045, 1053, 1059}
SKIP_FOR_NOW = {1009, 1018, 1027, 1036, 1045, 1054, 1055, 1056, 1059, 1060}
TOP_CONTENT = {1001, 1002, 1003, 1007, 1008, 1011, 1017, 1024, 1030, 1041, 1048, 1058}
AUDIT_IDS = {1001, 1002, 1003, 1004, 1007, 1008, 1015, 1022, 1031, 1040, 1050, 1060}


def build_records():
    tiers = ["S", "A", "A", "B", "B", "B", "C", "C", "D", "A"]
    priorities = [
        "1 — Essential", "2 — High", "2 — High", "3 — Medium", "3 — Medium",
        "4 — Low", "5 — Skip for Now", "3 — Medium", "4 — Low", "1 — Essential",
    ]
    rows = []
    for index, (title, year, director, genre) in enumerate(TITLES, start=1):
        channel = 100 - ((index * 7) % 41)
        taste = 98 - ((index * 11) % 43)
        rev06 = max(40, channel - ((index % 5) - 2))
        rows.append({
            "Rank": index,
            "Master ID": 1000 + index,
            "Title": title,
            "Year": year,
            "Director / Creator": director,
            "Genre": genre,
            "Rev06 Channel": rev06,
            "Rev07 Channel": channel,
            "Rev07 Taste Fit": taste,
            "Rev07 Tier": tiers[(index - 1) % len(tiers)],
            "Rev07 Viewing Priority": priorities[(index - 1) % len(priorities)],
            "Score Delta": channel - rev06,
            "Analytical Note": f"Synthetic fixture record {index}; fictional metadata for public demonstration.",
            "Rev07 Calibration Reason": "Synthetic calibration field used to exercise importer and reporting behavior.",
        })
    overrides = {
        1001: (100, 95, "S", "1 — Essential"),
        1002: (94, 89, "A", "2 — High"),
        1003: (90, 84, "A", "2 — High"),
        1004: (88, 81, "B", "3 — Medium"),
        1005: (86, 80, "B", "3 — Medium"),
        1006: (87, 82, "B", "3 — Medium"),
        1007: (91, 85, "A", "2 — High"),
        1008: (93, 88, "A", "2 — High"),
    }
    for row in rows:
        if row["Master ID"] in overrides:
            channel, taste, tier, priority = overrides[row["Master ID"]]
            row["Rev07 Channel"] = channel
            row["Rev07 Taste Fit"] = taste
            row["Rev07 Tier"] = tier
            row["Rev07 Viewing Priority"] = priority
            row["Score Delta"] = channel - row["Rev06 Channel"]
    return rows


def style_sheet(sheet):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        letter = column[0].column_letter
        header = str(column[0].value)
        width = 12
        if header in {"Title", "Director / Creator", "Director"}:
            width = 22
        elif header in {"Genre", "Rev07 Viewing Priority", "Viewing Priority"}:
            width = 20
        elif header in {"Analytical Note", "Rev07 Calibration Reason", "Why / Note", "Possible Episode Angle", "Central Question"}:
            width = 36
        elif header in {"Key Theme", "Rev06 Reason"}:
            width = 26
        sheet.column_dimensions[letter].width = width
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def append_sheet(workbook, name, headers, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)


def flag_rows(records_by_id, ids, label):
    rows = []
    for master_id in sorted(ids):
        row = records_by_id[master_id]
        rows.append([
            master_id, row["Title"], row["Year"], row["Director / Creator"], row["Genre"],
            row["Rev07 Channel"], row["Rev07 Taste Fit"], row["Rev07 Tier"],
            row["Rev07 Viewing Priority"], f"Synthetic {label} flag for public demo coverage.",
        ])
    return rows


def generate(output: Path):
    records = build_records()
    assert len(records) == 60
    by_id = {row["Master ID"]: row for row in records}

    workbook = Workbook()
    workbook.remove(workbook.active)

    append_sheet(
        workbook, "Rev07_Channel_Value_Ranking", MAIN_HEADERS,
        [[row[h] for h in MAIN_HEADERS] for row in records],
    )

    taste = sorted(records, key=lambda row: (-row["Rev07 Taste Fit"], row["Master ID"]))
    append_sheet(
        workbook, "Rev07_Taste_Fit_Ranking", MAIN_HEADERS,
        [[rank if h == "Rank" else row[h] for h in MAIN_HEADERS] for rank, row in enumerate(taste, 1)],
    )

    order = {"1 — Essential": 1, "2 — High": 2, "3 — Medium": 3, "4 — Low": 4, "5 — Skip for Now": 5}
    priority = sorted(records, key=lambda row: (order[row["Rev07 Viewing Priority"]], -row["Rev07 Channel"], row["Master ID"]))
    append_sheet(
        workbook, "Rev07_Viewing_Priority_Ranking", MAIN_HEADERS,
        [[rank if h == "Rank" else row[h] for h in MAIN_HEADERS] for rank, row in enumerate(priority, 1)],
    )

    append_sheet(workbook, "Rev07_Not_For_Channel", FLAG_HEADERS, flag_rows(by_id, NOT_FOR_CHANNEL, "not-for-channel"))
    append_sheet(workbook, "Rev07_Not_For_My_Taste", FLAG_HEADERS, flag_rows(by_id, NOT_FOR_TASTE, "not-for-taste"))
    append_sheet(workbook, "Rev07_Skip_For_Now", FLAG_HEADERS, flag_rows(by_id, SKIP_FOR_NOW, "skip-for-now"))

    top_rows = []
    for rank, master_id in enumerate(sorted(TOP_CONTENT), 1):
        row = by_id[master_id]
        top_rows.append([
            rank, master_id, row["Title"], row["Year"], row["Director / Creator"],
            row["Rev07 Channel"], row["Rev07 Taste Fit"], row["Rev07 Tier"],
            "Decision under uncertainty", "What changes when evidence is incomplete?",
            "Use the fictional scenario to test conflicting interpretations.",
            "Synthetic candidate; no relationship to the private archive.",
        ])
    append_sheet(workbook, "Rev07_Top_Content_Candidates", TOP_HEADERS, top_rows)

    audit_rows = []
    for master_id in sorted(AUDIT_IDS):
        row = by_id[master_id]
        previous_taste = max(0, row["Rev07 Taste Fit"] - 2)
        previous_tier = "A" if row["Rev07 Tier"] == "S" else row["Rev07 Tier"]
        previous_priority = "3 — Medium" if row["Rev07 Viewing Priority"] == "2 — High" else row["Rev07 Viewing Priority"]
        audit_rows.append([
            master_id, row["Title"], row["Year"], row["Director / Creator"],
            row["Rev06 Channel"], row["Rev07 Channel"], row["Score Delta"],
            previous_taste, row["Rev07 Taste Fit"], previous_tier, row["Rev07 Tier"],
            previous_priority, row["Rev07 Viewing Priority"],
            "Synthetic previous calibration state.", row["Rev07 Calibration Reason"],
        ])
    append_sheet(workbook, "Rev07_Score_Audit", AUDIT_HEADERS, audit_rows)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


if __name__ == "__main__":
    parser = ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "sample-data" / "Rev07_Sample_Catalog_60.xlsx",
    )
    args = parser.parse_args()
    generate(args.output)
    print(args.output)
