import hashlib
from pathlib import Path

from openpyxl import load_workbook

from master_film_curator.catalog.models import CatalogRecord
from master_film_curator.db.sqlite import event

MAIN = "Rev07_Channel_Value_Ranking"
FIELDS = {
    "master_id": "Master ID",
    "title": "Title",
    "year": "Year",
    "director": "Director / Creator",
    "genre": "Genre",
    "channel_score": "Rev07 Channel",
    "taste_fit": "Rev07 Taste Fit",
    "tier": "Rev07 Tier",
    "viewing_priority": "Rev07 Viewing Priority",
    "analytical_note": "Analytical Note",
    "calibration_reason": "Rev07 Calibration Reason",
}
FLAGS = {
    "not_for_channel": "Rev07_Not_For_Channel",
    "not_for_my_taste": "Rev07_Not_For_My_Taste",
    "skip_for_now": "Rev07_Skip_For_Now",
    "top_content_candidate": "Rev07_Top_Content_Candidates",
}


def sheet_rows(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    if len(headers) != len(set(headers)):
        raise ValueError(f"Duplicate headers in {sheet.title}")
    return [dict(zip(headers, row)) for row in iterator if any(v is not None for v in row)]


def read_catalog(path: Path, expected_count: int | None = 570) -> list[CatalogRecord]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = {MAIN, *FLAGS.values()} - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(f"Missing required sheets: {sorted(missing_sheets)}")
        rows = sheet_rows(workbook[MAIN])
        if not rows or (set(FIELDS.values()) - set(rows[0])):
            raise ValueError("Missing required catalog columns or empty catalog")
        if expected_count is not None and len(rows) != expected_count:
            raise ValueError(f"Expected {expected_count} catalog records; found {len(rows)}")
        records = []
        for row in rows:
            values = {field: row[column] for field, column in FIELDS.items()}
            if str(values["viewing_priority"]).split(" ")[0] not in {"1", "2", "3", "4", "5"}:
                raise ValueError(f"Invalid viewing priority: {values['viewing_priority']}")
            hint = f"{values['title']} {values['genre']}".casefold()
            records.append(
                CatalogRecord(
                    **values,
                    raw=row,
                    is_series=any(s in hint for s in ["miniseries", "tv series", "reality-tv", "tv-series"]),
                )
            )
        by_id = {r.master_id: r for r in records}
        if len(by_id) != len(records):
            raise ValueError("Duplicate Master IDs")
        for flag, sheet in FLAGS.items():
            seen = set()
            for row in sheet_rows(workbook[sheet]):
                mid = row["Master ID"]
                if mid not in by_id or mid in seen:
                    raise ValueError(f"Orphan or duplicate Master ID {mid} in {sheet}")
                seen.add(mid)
                setattr(by_id[mid], flag, True)
        return records
    finally:
        workbook.close()


def import_catalog(db, path: Path, expected_count=570) -> int:
    records = read_catalog(path, expected_count)
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    old = db.execute("SELECT value FROM metadata WHERE key='catalog_sha256'").fetchone()
    if old and old[0] == checksum:
        return len(records)
    old_ids = {r[0] for r in db.execute("SELECT master_id FROM catalog")}
    new_ids = {r.master_id for r in records}
    if old_ids - new_ids:
        raise ValueError("Catalog removed IDs; use a new database to preserve existing audit history")
    for r in records:
        db.execute(
            "INSERT INTO catalog VALUES(?,?) ON CONFLICT(master_id) DO UPDATE SET data=excluded.data",
            (r.master_id, r.model_dump_json()),
        )
    db.execute("INSERT OR REPLACE INTO metadata VALUES('catalog_sha256',?)", (checksum,))
    db.execute("INSERT OR REPLACE INTO metadata VALUES('catalog_source',?)", (str(path),))
    db.execute("UPDATE plans SET status='stale' WHERE status='planned'")
    # Scores and tiers are always imported verbatim. Automated matches must be recomputed.
    db.execute("UPDATE media SET master_id=NULL,status='unmatched',confidence=0,candidates='[]'")
    event(db, "catalog_import", {"count": len(records), "sha256": checksum, "source": str(path)})
    return len(records)
